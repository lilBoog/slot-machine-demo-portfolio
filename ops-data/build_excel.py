"""從 jinli.db 產生兩份 Excel：
  金鯉躍門_日報_2026-09.xlsx   每日營運報表模板（原始資料 → 公式計算 → 週報、版本比較、圖表）
  金鯉躍門_v1.3.1_驗收.xlsx    驗收文件（測試案例、設定核對、數學驗收、Bug 清單、回歸測試）
所有計算欄位都是 Excel 公式，換掉原始資料就會重算。
"""
import datetime as dt
import os
import sqlite3

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from slot_math import BET_LINES, DEFAULT, Machine, analytic, with_changes

HERE = os.path.dirname(os.path.abspath(__file__))
con = sqlite3.connect(os.path.join(HERE, "jinli.db"))

FONT = "Microsoft JhengHei"   # 中文內容需要中文字型
F_BODY = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")      # 藍字：輸入值
F_LINK = Font(name=FONT, size=10, color="008000")       # 綠字：引用其他工作表
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_NOTE = Font(name=FONT, size=9, color="595959")
FILL_HEAD = PatternFill("solid", fgColor="7A2418")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_SUB = PatternFill("solid", fgColor="F2F2F2")
RED_FILL = PatternFill("solid", fgColor="F8D7D3")
GREEN_FILL = PatternFill("solid", fgColor="DDF0E2")
AMBER_FILL = PatternFill("solid", fgColor="FCEBC5")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")

PCT = "0.00%"
INT = "#,##0"
DEC2 = "0.00"
DATE = "yyyy-mm-dd"


def header(ws, row, titles, widths=None):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, Alignment(horizontal="center", vertical="center", wrap_text=True), BOX
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 30


def put(ws, row, col, value, font=F_BODY, fmt=None, fill=None, align=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font, cell.border = font, BOX
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    return cell


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_NOTE


# =====================================================================
# 一、日報
# =====================================================================
wb = Workbook()

# ---------- 使用說明 ----------
ws = wb.active
ws.title = "使用說明"
title(ws, "金鯉躍門 每日營運報表（2026 年 9 月）", "示範資料：由 build_data.py 依遊戲數學模型模擬生成，非實際營運數據")
guide = [
    ("每日操作步驟", ""),
    ("1", "在 jinli.db 執行 queries.sql 的 [Q01]–[Q03]，確認沒有重複紀錄、錯誤數值或缺漏日期"),
    ("2", "執行 [Q04]，把結果貼到「原始資料」工作表（從 A2 開始，欄位順序相同）"),
    ("3", "「日報」「週報」「版本比較」「圖表」會自動重算，不需要手動修改"),
    ("4", "檢查「日報」狀態欄：標紅的「警示」要逐筆確認原因，並在備註欄說明"),
    ("5", "與財務或後台的總投注核對一致後，寄出或更新共用檔案"),
    ("", ""),
    ("顏色說明", ""),
    ("黃底藍字", "輸入值：原始資料、遊戲設定、警示門檻。只有這些儲存格需要手動修改"),
    ("黑字", "公式計算，請勿直接修改"),
    ("綠字", "引用其他工作表的值"),
    ("", ""),
    ("欄位定義", ""),
    ("實際 RTP", "總派彩 ÷ 總投注。派彩包含該轉觸發的免費遊戲贏分"),
    ("等效轉數", "總投注 ÷ 20。押注大小不同時，換算成「每線 1 點」的轉數，用來計算標準誤"),
    ("標準誤", "σ ÷ √等效轉數。σ 為 PAR 表的單局標準差（× 總押注）"),
    ("z 值", "(實際 RTP − 理論 RTP) ÷ 標準誤。|z| 超過門檻即警示"),
    ("觸發偏差 z", "(實際觸發次數 − 期望觸發次數) ÷ √期望次數。觸發次數近似卜瓦松分布"),
    ("累計 RTP", "月初至當日的總派彩 ÷ 總投注。轉數越多越接近理論值"),
    ("", ""),
    ("警示門檻為什麼用 3", "每天檢查一次，門檻 2 時單日誤報率 4.6%，30 天內至少誤報一次的機率約 75%；門檻 3 時約 7.8%"),
]
for i, (a, b) in enumerate(guide, 4):
    ws.cell(row=i, column=1, value=a).font = F_BOLD if b == "" or a in ("黃底藍字", "黑字", "綠字") else F_BODY
    ws.cell(row=i, column=2, value=b).font = F_BODY
    ws.cell(row=i, column=2).alignment = WRAP
ws["A12"].fill, ws["A12"].font = FILL_INPUT, F_INPUT
ws["A14"].font = F_LINK
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 95

# ---------- 遊戲設定 ----------
ws = wb.create_sheet("遊戲設定")
title(ws, "遊戲設定（來源：games 資料表／PAR 表）")
header(ws, 4, ["版本", "理論 RTP", "單局標準差 σ", "觸發率", "平均免費場次", "上線日", "結束日", "說明"], [10, 12, 14, 12, 14, 12, 12, 46])
games = con.execute("SELECT version, theoretical_rtp, sigma, trigger_rate, avg_fg_spins, release_date, end_date, note FROM games ORDER BY release_date").fetchall()
for r, g in enumerate(games, 5):
    put(ws, r, 1, g[0], F_INPUT, fill=FILL_INPUT)
    put(ws, r, 2, round(g[1], 6), F_INPUT, PCT, FILL_INPUT)
    put(ws, r, 3, round(g[2], 4), F_INPUT, "0.0000", FILL_INPUT)
    put(ws, r, 4, round(g[3], 6), F_INPUT, "0.000%", FILL_INPUT)
    put(ws, r, 5, round(g[4], 4), F_INPUT, DEC2, FILL_INPUT)
    put(ws, r, 6, dt.date.fromisoformat(g[5]), F_INPUT, DATE, FILL_INPUT)
    put(ws, r, 7, dt.date.fromisoformat(g[6]), F_INPUT, DATE, FILL_INPUT)
    put(ws, r, 8, g[7], F_INPUT, fill=FILL_INPUT)
put(ws, 10, 1, "z 警示門檻", F_BOLD)
put(ws, 10, 2, 3, F_INPUT, "0.0", FILL_INPUT)
ws["B10"].comment = Comment("每日檢查時建議用 3；改成 2 會大幅增加誤報", "報表")
put(ws, 11, 1, "固定區間 ±", F_BOLD)
put(ws, 11, 2, 0.03, F_INPUT, "0.0%", FILL_INPUT)
ws["C11"] = "對照用的簡單規則：實際 RTP 偏離理論值超過此幅度就警示"
ws["C11"].font = F_NOTE
put(ws, 12, 1, "等效轉數基準", F_BOLD)
put(ws, 12, 2, 20, F_INPUT, INT, FILL_INPUT)
ws["C12"] = "每線 1 點 × 20 線 = 每轉 20 點"
ws["C12"].font = F_NOTE
ws["A13"] = "理論 RTP、σ、觸發率依 PAR 表規格；v1.3 的規格與 v1.2 相同，實際上線版本的設定錯誤不會反映在理論值上，這正是要靠報表抓出來的差異。"
ws["A13"].font = F_NOTE
ws.freeze_panes = "A5"

# ---------- 原始資料 ----------
ws = wb.create_sheet("原始資料")
raw_cols = ["日期", "版本", "轉數", "玩家數", "總投注", "總派彩", "觸發次數", "最高單局倍數"]
header(ws, 1, raw_cols, [12, 10, 10, 10, 13, 13, 11, 13])
daily = con.execute("""
  SELECT date(spin_time), version, COUNT(*), COUNT(DISTINCT player_id), SUM(bet), SUM(win), SUM(fg_triggered),
         ROUND(MAX(win * 1.0 / bet), 1)
  FROM v_spins GROUP BY date(spin_time), version ORDER BY 1""").fetchall()
for r, d in enumerate(daily, 2):
    vals = [dt.date.fromisoformat(d[0])] + list(d[1:])
    fmts = [DATE, None, INT, INT, INT, INT, INT, "0.0"]
    for c, (v, f) in enumerate(zip(vals, fmts), 1):
        put(ws, r, c, v, F_INPUT, f, FILL_INPUT)
ws.freeze_panes = "A2"
N = len(daily)
LAST = N + 1

# ---------- 日報 ----------
ws = wb.create_sheet("日報")
cols = ["日期", "版本", "轉數", "玩家數", "總投注", "總派彩", "莊家淨收入", "實際 RTP", "理論 RTP", "σ",
        "等效轉數", "標準誤", "z 值", "p 值", "觸發次數", "期望觸發", "觸發偏差 z", "最高單局倍數",
        "累計 RTP", "z 檢定", "固定區間", "警示上界", "警示下界", "備註"]
header(ws, 1, cols, [11, 8, 9, 8, 12, 12, 12, 10, 10, 8, 10, 9, 8, 8, 9, 9, 10, 11, 10, 8, 9, 10, 10, 36])
GS = "遊戲設定"
for r in range(2, LAST + 1):
    f = {
        1: (f"=原始資料!A{r}", DATE, F_LINK), 2: (f"=原始資料!B{r}", None, F_LINK), 3: (f"=原始資料!C{r}", INT, F_LINK),
        4: (f"=原始資料!D{r}", INT, F_LINK), 5: (f"=原始資料!E{r}", INT, F_LINK), 6: (f"=原始資料!F{r}", INT, F_LINK),
        7: (f"=E{r}-F{r}", INT, F_BODY),
        8: (f"=IFERROR(F{r}/E{r},0)", PCT, F_BODY),
        9: (f"=INDEX({GS}!$B$5:$B$9,MATCH(B{r},{GS}!$A$5:$A$9,0))", PCT, F_LINK),
        10: (f"=INDEX({GS}!$C$5:$C$9,MATCH(B{r},{GS}!$A$5:$A$9,0))", "0.000", F_LINK),
        11: (f"=E{r}/{GS}!$B$12", INT, F_BODY),
        12: (f"=J{r}/SQRT(K{r})", "0.00%", F_BODY),
        13: (f"=(H{r}-I{r})/L{r}", DEC2, F_BODY),
        14: (f"=2*(1-NORMSDIST(ABS(M{r})))", "0.0000", F_BODY),
        15: (f"=原始資料!G{r}", INT, F_LINK),
        16: (f"=C{r}*INDEX({GS}!$D$5:$D$9,MATCH(B{r},{GS}!$A$5:$A$9,0))", "0.0", F_BODY),
        17: (f"=(O{r}-P{r})/SQRT(P{r})", DEC2, F_BODY),
        18: (f"=原始資料!H{r}", "0.0", F_LINK),
        19: (f"=SUM($F$2:F{r})/SUM($E$2:E{r})", PCT, F_BODY),
        20: (f'=IF(ABS(M{r})>{GS}!$B$10,"警示","正常")', None, F_BODY),
        21: (f'=IF(ABS(H{r}-I{r})>{GS}!$B$11,"警示","正常")', None, F_BODY),
        22: (f"=I{r}+{GS}!$B$10*L{r}", PCT, F_BODY),
        23: (f"=I{r}-{GS}!$B$10*L{r}", PCT, F_BODY),
    }
    for c, (v, fmt, font) in f.items():
        put(ws, r, c, v, font, fmt)
    put(ws, r, 24, None)
notes = {"2026-09-17": "正常版本的誤報（|z| 剛好超過 3）；另有 2 筆押注為 0 的錯誤紀錄已由 v_spins 排除",
         "2026-09-21": "v1.3 上線首日即警示 → 以 SQL [Q06][Q07] 查出免費遊戲倍數誤設為 ×3",
         "2026-09-24": "v1.3 最後一日，9/25 上線 v1.3.1 修正",
         "2026-09-12": "原始資料有 3 筆重複寫入，已由 v_spins 排除"}
for r, d in enumerate(daily, 2):
    if d[0] in notes:
        ws.cell(row=r, column=24, value=notes[d[0]]).font = F_BODY
T = LAST + 1
put(ws, T, 1, "月合計", F_BOLD, fill=FILL_SUB)
for c in range(2, 25):
    put(ws, T, c, None, fill=FILL_SUB)
for c, col in ((3, "C"), (5, "E"), (6, "F"), (7, "G"), (15, "O")):
    put(ws, T, c, f"=SUM({col}2:{col}{LAST})", F_BOLD, INT, FILL_SUB)
put(ws, T, 8, f"=F{T}/E{T}", F_BOLD, PCT, FILL_SUB)
put(ws, T, 16, f"=SUM(P2:P{LAST})", F_BOLD, "0.0", FILL_SUB)
put(ws, T, 20, f'=COUNTIF(T2:T{LAST},"警示")&" 天警示"', F_BOLD, None, FILL_SUB)
put(ws, T, 21, f'=COUNTIF(U2:U{LAST},"警示")&" 天警示"', F_BOLD, None, FILL_SUB)
rng = f"A2:X{LAST}"
ws.conditional_formatting.add(rng, FormulaRule(formula=['$T2="警示"'], fill=RED_FILL))
ws.conditional_formatting.add(f"U2:U{LAST}", FormulaRule(formula=['$U2="警示"'], font=Font(name=FONT, color="9C0006", bold=True)))
ws.freeze_panes = "C2"

# ---------- 週報 ----------
ws = wb.create_sheet("週報")
title(ws, "週報（依「日報」自動彙總）")
header(ws, 3, ["週別", "起日", "迄日", "轉數", "總投注", "總派彩", "莊家淨收入", "實際 RTP", "z 檢定警示天數", "固定區間警示天數"],
       [10, 12, 12, 11, 13, 13, 13, 11, 14, 15])
weeks = [("第 1 週", "2026-09-01", "2026-09-06"), ("第 2 週", "2026-09-07", "2026-09-13"), ("第 3 週", "2026-09-14", "2026-09-20"),
         ("第 4 週", "2026-09-21", "2026-09-27"), ("第 5 週", "2026-09-28", "2026-09-30")]
D = f"日報!$A$2:$A${LAST}"
for r, (w, a, b) in enumerate(weeks, 4):
    put(ws, r, 1, w)
    put(ws, r, 2, dt.date.fromisoformat(a), F_INPUT, DATE, FILL_INPUT)
    put(ws, r, 3, dt.date.fromisoformat(b), F_INPUT, DATE, FILL_INPUT)
    for c, col in ((4, "C"), (5, "E"), (6, "F"), (7, "G")):
        put(ws, r, c, f'=SUMIFS(日報!${col}$2:${col}${LAST},{D},">="&$B{r},{D},"<="&$C{r})', F_BODY, INT)
    put(ws, r, 8, f"=IFERROR(F{r}/E{r},0)", F_BODY, PCT)
    put(ws, r, 9, f'=COUNTIFS({D},">="&$B{r},{D},"<="&$C{r},日報!$T$2:$T${LAST},"警示")', F_BODY, INT)
    put(ws, r, 10, f'=COUNTIFS({D},">="&$B{r},{D},"<="&$C{r},日報!$U$2:$U${LAST},"警示")', F_BODY, INT)
WE = 3 + len(weeks)
put(ws, WE + 1, 1, "合計", F_BOLD, fill=FILL_SUB)
for c in (2, 3):
    put(ws, WE + 1, c, None, fill=FILL_SUB)
for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (9, "I"), (10, "J")):
    put(ws, WE + 1, c, f"=SUM({col}4:{col}{WE})", F_BOLD, INT, FILL_SUB)
put(ws, WE + 1, 8, f"=F{WE + 1}/E{WE + 1}", F_BOLD, PCT, FILL_SUB)
ws.cell(row=WE + 3, column=1, value="週別起迄日為輸入值（黃底），新增週別時複製最後一列並修改日期即可。").font = F_NOTE

# ---------- 版本比較 ----------
ws = wb.create_sheet("版本比較")
title(ws, "版本比較（各版本上線期間的累計表現）")
header(ws, 3, ["版本", "天數", "轉數", "總投注", "總派彩", "實際 RTP", "理論 RTP", "σ", "等效轉數", "z 值", "判定", "說明"],
       [10, 7, 11, 13, 13, 11, 11, 8, 11, 8, 10, 46])
V = f"日報!$B$2:$B${LAST}"
for i, g in enumerate(games):
    r = 4 + i
    put(ws, r, 1, f"={GS}!A{5 + i}", F_LINK)
    put(ws, r, 2, f"=COUNTIF({V},A{r})", F_BODY, INT)
    for c, col in ((3, "C"), (4, "E"), (5, "F")):
        put(ws, r, c, f"=SUMIF({V},A{r},日報!${col}$2:${col}${LAST})", F_BODY, INT)
    put(ws, r, 6, f"=IFERROR(E{r}/D{r},0)", F_BODY, PCT)
    put(ws, r, 7, f"={GS}!B{5 + i}", F_LINK, PCT)
    put(ws, r, 8, f"={GS}!C{5 + i}", F_LINK, "0.000")
    put(ws, r, 9, f"=D{r}/{GS}!$B$12", F_BODY, INT)
    put(ws, r, 10, f"=(F{r}-G{r})/(H{r}/SQRT(I{r}))", F_BODY, DEC2)
    put(ws, r, 11, f'=IF(ABS(J{r})>{GS}!$B$10,"異常","正常")', F_BODY)
    put(ws, r, 12, f"={GS}!H{5 + i}", F_LINK)
ws.conditional_formatting.add("A4:L6", FormulaRule(formula=['$K4="異常"'], fill=RED_FILL))
ws.cell(row=8, column=1, value="v1.3 累計 RTP 明顯偏高，與免費遊戲倍數設定錯誤相符；v1.3.1 修正後回到正常範圍。").font = F_NOTE

# ---------- 圖表 ----------
ws = wb.create_sheet("圖表")
title(ws, "圖表（資料來源：日報、週報）")
dr = wb["日報"]
ch = LineChart()
ch.title = "每日實際 RTP 與 z 檢定警示範圍"
ch.y_axis.title = "RTP"
ch.y_axis.number_format = "0%"
ch.height, ch.width = 9, 26
for col in (8, 9, 22, 23):
    ch.add_data(Reference(dr, min_col=col, min_row=1, max_row=LAST), titles_from_data=True)
ch.set_categories(Reference(dr, min_col=1, min_row=2, max_row=LAST))
ch.x_axis.number_format = "mm/dd"
ch.series[0].graphicalProperties.line.solidFill = "A82C1C"
ch.series[0].graphicalProperties.line.width = 28000
ch.series[1].graphicalProperties.line.solidFill = "333333"
for s in ch.series[2:]:
    s.graphicalProperties.line.solidFill = "9AA2AC"
    s.graphicalProperties.line.dashStyle = "dash"
ws.add_chart(ch, "A3")
ch2 = LineChart()
ch2.title = "累計 RTP"
ch2.y_axis.number_format = "0%"
ch2.height, ch2.width = 9, 26
ch2.add_data(Reference(dr, min_col=19, min_row=1, max_row=LAST), titles_from_data=True)
ch2.add_data(Reference(dr, min_col=9, min_row=1, max_row=LAST), titles_from_data=True)
ch2.set_categories(Reference(dr, min_col=1, min_row=2, max_row=LAST))
ch2.x_axis.number_format = "mm/dd"
ch2.series[0].graphicalProperties.line.solidFill = "2D5FA6"
ch2.series[1].graphicalProperties.line.solidFill = "333333"
ws.add_chart(ch2, "A23")
wr = wb["週報"]
ch3 = BarChart()
ch3.title = "每週莊家淨收入（點）"
ch3.height, ch3.width = 9, 26
ch3.add_data(Reference(wr, min_col=7, min_row=3, max_row=WE), titles_from_data=True)
ch3.set_categories(Reference(wr, min_col=1, min_row=4, max_row=WE))
ch3.series[0].graphicalProperties.solidFill = "A0501A"
ws.add_chart(ch3, "A43")

for s in wb.worksheets:
    s.sheet_view.showGridLines = s.title not in ("使用說明", "圖表")
wb.save(os.path.join(HERE, "金鯉躍門_日報_2026-09.xlsx"))

# =====================================================================
# 二、驗收
# =====================================================================
wb = Workbook()

# ---------- 摘要 ----------
ws = wb.active
ws.title = "摘要"
title(ws, "金鯉躍門 v1.3.1 驗收摘要", "示範文件：情境為 v1.3 上線後被報表發現設定錯誤，修正版 v1.3.1 的驗收紀錄")
info = [("遊戲", "金鯉躍門"), ("驗收版本", "v1.3.1"), ("驗收期間", "2026-09-24 ～ 2026-09-25"), ("驗收人員", "（填入姓名）"),
        ("驗收依據", "遊戲規格書 v1.3、PAR 表 2026-08、Bug 清單 BUG-01")]
for i, (k, v) in enumerate(info, 4):
    put(ws, i, 1, k, F_BOLD, fill=FILL_SUB)
    put(ws, i, 2, v, F_INPUT, fill=FILL_INPUT)
put(ws, 10, 1, "測試案例統計", F_BOLD)
stats = [("案例總數", '=COUNTIF(測試案例!A2:A200,"TC-*")'), ("通過", '=COUNTIF(測試案例!H2:H200,"通過")'), ("不通過", '=COUNTIF(測試案例!H2:H200,"不通過")'),
         ("阻擋（無法測試）", '=COUNTIF(測試案例!H2:H200,"阻擋")'), ("未執行", '=COUNTIF(測試案例!H2:H200,"未執行")'),
         ("通過率（已執行）", "=IFERROR(B12/(B11-B15),0)")]
for i, (k, v) in enumerate(stats, 11):
    put(ws, i, 1, k, F_BODY)
    put(ws, i, 2, v, F_BODY, PCT if "率" in k else INT)
put(ws, 18, 1, "上線條件", F_BOLD)
cond = [("高嚴重度未關閉 Bug", '=COUNTIFS(Bug清單!D2:D50,"高",Bug清單!J2:J50,"<>已關閉")', INT),
        ("中低嚴重度未關閉 Bug", '=COUNTIFS(Bug清單!D2:D50,"中",Bug清單!J2:J50,"<>已關閉")+COUNTIFS(Bug清單!D2:D50,"低",Bug清單!J2:J50,"<>已關閉")', INT),
        ("設定核對不一致（v1.3.1）", '=COUNTIF(設定核對!G2:G200,"不一致")', INT),
        ("數學驗收（RTP）", "=數學驗收!C16", None),
        ("數學驗收（觸發率）", "=數學驗收!C20", None)]
for i, (k, v, fmt) in enumerate(cond, 19):
    put(ws, i, 1, k, F_BODY)
    put(ws, i, 2, v, F_LINK if "!" in v and "COUNT" not in v else F_BODY, fmt)
put(ws, 25, 1, "驗收結論", F_BOLD, fill=FILL_SUB)
put(ws, 25, 2, '=IF(AND(B19=0,B21=0,B22="通過",B23="通過"),IF(B20=0,"可上線","可上線（"&B20&" 項中低嚴重度問題列入後續修正）"),"不可上線")', F_BOLD, fill=FILL_SUB)
ws.conditional_formatting.add("B25", FormulaRule(formula=['LEFT($B$25,3)="可上線"'], fill=GREEN_FILL))
ws.conditional_formatting.add("B25", FormulaRule(formula=['$B$25="不可上線"'], fill=RED_FILL))
ws["A27"] = "結論規則：高嚴重度 Bug 全部關閉、設定與 PAR 表完全一致、數學驗收通過，才可上線。"
ws["A27"].font = F_NOTE
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 60

# ---------- 測試案例 ----------
ws = wb.create_sheet("測試案例")
tc_cols = ["編號", "類別", "測試項目", "前置條件", "步驟／盤面", "預期結果", "實際結果", "結果", "嚴重度", "Bug 編號", "備註"]
header(ws, 1, tc_cols, [8, 10, 24, 20, 30, 28, 22, 8, 7, 9, 28])
P = "通過"
cases = [
    ("設定核對", "輪帶停點數與 PAR 表一致", "取得 v1.3.1 匯出設定", "依「設定核對」頁逐項比對 50 項", "全部一致", "全部一致", P, "", "", ""),
    ("設定核對", "賠率表與 PAR 表一致", "同上", "逐項比對 27 項賠率", "全部一致", "全部一致", P, "", "", ""),
    ("設定核對", "分散賠付與免費遊戲設定一致", "同上", "比對分散賠付 3 項、免費場次、免費倍數", "全部一致", "全部一致", P, "", "", "v1.3 的免費倍數為 3，見 BUG-01"),
    ("連線判定", "基本 3 連", "每線押 1，測試模式指定停輪", "第 1 線：寶 寶 寶 J Q", "寶 3 連，贏 50", "贏 50", P, "", "", ""),
    ("連線判定", "基本 4 連", "同上", "第 1 線：寶 寶 寶 寶 Q", "寶 4 連，贏 200", "贏 200", P, "", "", ""),
    ("連線判定", "基本 5 連", "同上", "第 1 線：寶 寶 寶 寶 寶", "寶 5 連，贏 1,000", "贏 1,000", P, "", "", ""),
    ("連線判定", "必須從最左軸起算", "同上", "第 1 線：J 寶 寶 寶 寶", "不中獎", "不中獎", P, "", "", "邊界：右側 4 連不算"),
    ("連線判定", "中間斷開不算連線", "同上", "第 1 線：寶 寶 J 寶 寶", "不中獎（只有 2 連）", "不中獎", P, "", "", ""),
    ("連線判定", "非直線連線", "同上", "第 4 線（V 形）：上 中 下 中 上 皆為 燈", "燈 5 連，贏 300", "贏 300", P, "", "", "抽查 20 條線中的 5 條"),
    ("百搭", "百搭替代延伸連線", "同上", "第 1 線：寶 鯉 鯉 鯉 鯉", "寶 5 連，贏 1,000", "贏 1,000", P, "", "", ""),
    ("百搭", "百搭不替代分散符號", "同上", "第 1 線：寶 鯉 門 寶 寶", "連線不中獎（寶 2 連）", "連線不中獎", P, "", "", "門另依分散規則計算"),
    ("百搭", "第 1 軸不會出現百搭", "測試模式逐一指定第 1 軸停點", "第 1 軸 36 個停點全部檢查", "沒有任何停點出現鯉", "未出現", P, "", "", "設計規則"),
    ("分散", "3 個門：賠付並觸發", "每線押 1", "第 1、3、5 軸各出現 1 個門", "贏 40（2 × 總押注 20），觸發 10 場", "贏 40，觸發 10 場", P, "", "", ""),
    ("分散", "4 個門", "同上", "第 1–4 軸各出現 1 個門", "贏 200，觸發 10 場", "贏 200，觸發 10 場", P, "", "", ""),
    ("分散", "5 個門", "同上", "五軸各出現 1 個門", "贏 1,000，觸發 10 場", "贏 1,000，觸發 10 場", P, "", "", ""),
    ("分散", "同一軸視窗最多 1 個門", "取得輪帶排列", "檢查每軸相鄰兩個門的間距", "間距皆 ≥ 3", "最小間距 18", P, "", "", ""),
    ("免費遊戲", "免費場次", "觸發免費遊戲", "計算實際進行場次（無再觸發）", "10 場", "10 場", P, "", "", ""),
    ("免費遊戲", "免費遊戲贏分倍數", "免費遊戲中，每線押 1", "第 1 線：寶 寶 寶 J Q", "贏 100（50 × 2）", "贏 100", P, "", "BUG-01", "v1.3 實際贏 150，回歸驗證已修正"),
    ("免費遊戲", "免費遊戲中再觸發", "免費遊戲第 4 場", "出現 3 個門", "剩餘場次 +10（6 → 16）", "6 → 16", P, "", "", ""),
    ("免費遊戲", "免費遊戲中分散賠付也乘倍數", "免費遊戲中", "出現 3 個門", "分散贏 80（40 × 2）", "贏 80", P, "", "BUG-01", ""),
    ("押注與餘額", "總押注扣款正確", "餘額 10,000", "每線押 1、2、5 各轉一次", "扣 20、40、100", "扣 20、40、100", P, "", "", ""),
    ("押注與餘額", "贏分依線注等比放大", "每線押 2", "第 1 線：寶 寶 寶 J Q", "贏 100（50 × 2）", "贏 100", P, "", "", ""),
    ("押注與餘額", "餘額不足不可轉動", "餘額 15，總押注 20", "按轉動", "不可轉動，提示餘額不足", "提示餘額不足", P, "", "", "邊界值"),
    ("押注與餘額", "餘額剛好等於押注", "餘額 20，總押注 20", "按轉動", "可轉動，轉後餘額 0（加贏分）", "符合", P, "", "", "邊界值"),
    ("介面", "贏分線位置正確亮起", "任一中獎盤面", "比對亮起位置與中獎線", "位置一致", "一致", P, "", "", "抽查 10 次"),
    ("介面", "餘額顯示千分位", "餘額 10,000", "查看餘額欄", "顯示 10,000", "顯示 10000", "不通過", "低", "BUG-02", ""),
    ("異常處理", "免費遊戲中斷線重連", "免費遊戲第 5 場", "中斷網路後重新連線", "回到第 5 場，累計贏分不變", "", "阻擋", "", "", "測試環境無法模擬斷線，已請開發提供工具"),
    ("異常處理", "快速連點轉動", "主遊戲", "1 秒內連點轉動 5 次", "只執行 1 次轉動", "只執行 1 次", P, "", "", ""),
    ("紀錄", "每轉寫入一筆紀錄", "轉動 100 次（含 1 次免費遊戲）", "以 SQL 查詢 spins 筆數與 win 欄位", "100 筆；win 含免費遊戲贏分", "符合", P, "", "", "見 queries.sql"),
    ("數學驗收", "RTP 落在理論值 ±3 標準誤", "模擬 100 萬轉", "見「數學驗收」頁", "|z| ≤ 3", "見數學驗收頁", P, "", "", "結果欄依數學驗收頁判定"),
    ("數學驗收", "免費遊戲觸發率與平均場次", "同上", "見「數學驗收」頁", "|z| ≤ 3、平均場次接近理論值", "見數學驗收頁", P, "", "", ""),
]
for i, c in enumerate(cases):
    r = 2 + i
    put(ws, r, 1, f"TC-{i + 1:03d}", F_BODY)
    for j, v in enumerate(c, 2):
        put(ws, r, j, v if v != "" else None, F_INPUT if j >= 7 else F_BODY, None, FILL_INPUT if j >= 7 else None, WRAP)
# 數學驗收兩列的結果由公式帶入
ws.cell(row=2 + len(cases) - 2, column=8, value="=數學驗收!C16").font = F_LINK
ws.cell(row=2 + len(cases) - 1, column=8, value="=數學驗收!C20").font = F_LINK
for rr in (2 + len(cases) - 2, 2 + len(cases) - 1):
    ws.cell(row=rr, column=8).fill = PatternFill(fill_type=None)
TC_LAST = 1 + len(cases)
dv1 = DataValidation(type="list", formula1='"通過,不通過,阻擋,未執行"', allow_blank=True, showErrorMessage=True, errorTitle="輸入錯誤", error="請從清單選擇：通過、不通過、阻擋、未執行")
dv2 = DataValidation(type="list", formula1='"高,中,低"', allow_blank=True, showErrorMessage=True, errorTitle="輸入錯誤", error="請從清單選擇：高、中、低")
ws.add_data_validation(dv1)
ws.add_data_validation(dv2)
dv1.add(f"H2:H{TC_LAST - 2}")
dv1.add(f"H{TC_LAST + 1}:H200")
dv2.add("I2:I200")
ws.conditional_formatting.add("H2:H200", FormulaRule(formula=['$H2="通過"'], fill=GREEN_FILL))
ws.conditional_formatting.add("H2:H200", FormulaRule(formula=['$H2="不通過"'], fill=RED_FILL))
ws.conditional_formatting.add("H2:H200", FormulaRule(formula=['$H2="阻擋"'], fill=AMBER_FILL))
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:K{TC_LAST}"
ws.cell(row=TC_LAST + 2, column=1, value="黃底藍字為測試人員填寫欄位；「結果」「嚴重度」只能從下拉清單選擇。新增案例時接續在最後一列。").font = F_NOTE

# ---------- 設定核對 ----------
ws = wb.create_sheet("設定核對")
header(ws, 1, ["參數", "說明", "PAR 表規格", "v1.3 實際", "v1.3.1 實際", "v1.3 判定", "v1.3.1 判定"], [16, 30, 11, 11, 11, 11, 11])
par = con.execute("SELECT param, description, value FROM config_par").fetchall()
build = {(v, p): val for v, p, val in con.execute("SELECT version, param, value FROM config_build")}
for i, (p, desc, val) in enumerate(par, 2):
    put(ws, i, 1, p)
    put(ws, i, 2, desc)
    put(ws, i, 3, val, F_INPUT, "0", FILL_INPUT)
    put(ws, i, 4, build[("v1.3", p)], F_INPUT, "0", FILL_INPUT)
    put(ws, i, 5, build[("v1.3.1", p)], F_INPUT, "0", FILL_INPUT)
    put(ws, i, 6, f'=IF(D{i}=C{i},"一致","不一致")')
    put(ws, i, 7, f'=IF(E{i}=C{i},"一致","不一致")')
PL = 1 + len(par)
ws.conditional_formatting.add(f"A2:G{PL}", FormulaRule(formula=['OR($F2="不一致",$G2="不一致")'], fill=RED_FILL))
put(ws, PL + 1, 1, "不一致項目數", F_BOLD, fill=FILL_SUB)
put(ws, PL + 1, 2, None, fill=FILL_SUB)
for c in (3, 4, 5):
    put(ws, PL + 1, c, None, fill=FILL_SUB)
put(ws, PL + 1, 6, f'=COUNTIF(F2:F{PL},"不一致")', F_BOLD, INT, FILL_SUB)
put(ws, PL + 1, 7, f'=COUNTIF(G2:G{PL},"不一致")', F_BOLD, INT, FILL_SUB)
ws.cell(row=PL + 3, column=1, value="資料來源：jinli.db 的 config_par（PAR 表）與 config_build（開發匯出的實際設定）；也可用 queries.sql [Q07] 直接列出不一致項目。").font = F_NOTE
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{PL}"

# ---------- 數學驗收 ----------
ws = wb.create_sheet("數學驗收")
title(ws, "數學驗收：100 萬轉模擬 vs PAR 表理論值")
spec = analytic(DEFAULT)
sig = float(con.execute("SELECT sigma FROM games WHERE version='v1.3.1'").fetchone()[0])
sims = {}
for v, cfg, seed in (("v1.3.1", DEFAULT, 31), ("v1.3（事後重跑）", with_changes(DEFAULT, mult=3), 13)):
    res = Machine(cfg).play(np.random.default_rng(seed), 1_000_000)
    sims[v] = {"n": 1_000_000, "win": int(res["win"].sum()), "trig": int(res["fg_triggered"].sum()), "fgs": int(res["fg_spins"].sum())}
header(ws, 3, ["項目", "說明", "v1.3.1", "v1.3（事後重跑）"], [24, 46, 16, 18])
rows = [
    ("模擬轉數", "每線押 1，總押注 20", "n", INT, True),
    ("總押注", "轉數 × 20", None, INT, False),
    ("總派彩", "含免費遊戲贏分", "win", INT, True),
    ("觸發次數", "主遊戲觸發免費遊戲的次數", "trig", INT, True),
    ("免費遊戲總場次", "含再觸發", "fgs", INT, True),
    ("理論 RTP", "PAR 表", "theo", PCT, True),
    ("單局標準差 σ", "PAR 表（× 總押注）", "sig", "0.0000", True),
    ("理論觸發率", "PAR 表", "pt", "0.000%", True),
    ("理論平均免費場次", "F ÷ (1 − F·p)", "efg", DEC2, True),
]
for i, (k, d, key, fmt, is_in) in enumerate(rows, 4):
    put(ws, i, 1, k, F_BOLD)
    put(ws, i, 2, d, F_NOTE)
    for j, v in enumerate(sims, 3):
        val = {"n": sims[v]["n"], "win": sims[v]["win"], "trig": sims[v]["trig"], "fgs": sims[v]["fgs"],
               "theo": round(spec["total"], 6), "sig": round(sig, 4), "pt": round(spec["pt"], 6), "efg": round(spec["e_fg"], 4)}.get(key)
        col = get_column_letter(j)
        if key is None:
            put(ws, i, j, f"={col}4*20", F_BODY, fmt)
        else:
            put(ws, i, j, val, F_INPUT, fmt, FILL_INPUT)
calc = [
    (13, "實際 RTP", "總派彩 ÷ 總投注", "={c}6/{c}5", PCT),
    (14, "標準誤", "σ ÷ √轉數", "={c}10/SQRT({c}4)", "0.000%"),
    (15, "RTP z 值", "(實際 − 理論) ÷ 標準誤", "=({c}13-{c}9)/{c}14", DEC2),
    (16, "RTP 判定", "|z| ≤ 3 為通過", '=IF(ABS({c}15)<=3,"通過","不通過")', None),
    (17, "95% 信賴區間", "實際 RTP ± 1.96 × 標準誤", '=TEXT({c}13-1.96*{c}14,"0.00%")&" ～ "&TEXT({c}13+1.96*{c}14,"0.00%")', None),
    (18, "實際觸發率", "觸發次數 ÷ 轉數", "={c}7/{c}4", "0.000%"),
    (19, "觸發次數 z 值", "(實際 − 期望) ÷ √期望（卜瓦松近似）", "=({c}7-{c}4*{c}11)/SQRT({c}4*{c}11)", DEC2),
    (20, "觸發判定", "|z| ≤ 3 為通過", '=IF(ABS({c}19)<=3,"通過","不通過")', None),
    (21, "實際平均免費場次", "免費總場次 ÷ 觸發次數", "={c}8/{c}7", DEC2),
]
for r, k, d, f, fmt in calc:
    put(ws, r, 1, k, F_BOLD)
    put(ws, r, 2, d, F_NOTE)
    for j in (3, 4):
        c = get_column_letter(j)
        put(ws, r, j, f.format(c=c), F_BOLD if r in (16, 20) else F_BODY, fmt)
ws.conditional_formatting.add("C16:D16", FormulaRule(formula=['C16="不通過"'], fill=RED_FILL))
ws.conditional_formatting.add("C16:D16", FormulaRule(formula=['C16="通過"'], fill=GREEN_FILL))
ws.conditional_formatting.add("C20:D20", FormulaRule(formula=['C20="不通過"'], fill=RED_FILL))
ws.conditional_formatting.add("C20:D20", FormulaRule(formula=['C20="通過"'], fill=GREEN_FILL))
ws["A23"] = "v1.3（事後重跑）：用 v1.3 的實際設定（免費倍數 ×3）重跑同樣的驗收，RTP 判定為不通過。代表若當時有做數學驗收，這個錯誤在上線前就能被擋下。"
ws["A23"].font = F_NOTE
ws["A24"] = "觸發率不受免費倍數影響，所以兩個版本的觸發判定都通過；問題只出現在 RTP。這跟營運資料 [Q06] 的發現一致。"
ws["A24"].font = F_NOTE

# ---------- Bug 清單 ----------
ws = wb.create_sheet("Bug清單")
header(ws, 1, ["編號", "標題", "發現版本", "嚴重度", "發現日期", "發現方式", "重現步驟", "預期結果", "實際結果", "狀態", "修正版本", "回歸結果"],
       [8, 26, 9, 7, 11, 26, 30, 18, 18, 8, 9, 22])
bugs = [
    ("BUG-01", "免費遊戲贏分倍數誤設為 ×3", "v1.3", "高", dt.date(2026, 9, 21), "日報 9/21 z 檢定警示 → SQL [Q06] 免費遊戲平均贏分偏高 → [Q07] 設定核對",
     "觸發免費遊戲後，第 1 線出現 寶 寶 寶 J Q", "贏 100（50 × 2）", "贏 150（50 × 3）", "已關閉", "v1.3.1", "9/25 回歸通過（TC-003、018、020、030）"),
    ("BUG-02", "餘額未顯示千分位", "v1.3.1", "低", dt.date(2026, 9, 25), "驗收 TC-026", "餘額 10,000 時查看餘額欄", "顯示 10,000", "顯示 10000", "待修", "", ""),
]
for i, b in enumerate(bugs, 2):
    for j, v in enumerate(b, 1):
        put(ws, i, j, v, F_INPUT, DATE if isinstance(v, dt.date) else None, FILL_INPUT, WRAP)
dv3 = DataValidation(type="list", formula1='"高,中,低"', allow_blank=True)
dv4 = DataValidation(type="list", formula1='"待修,修正中,待驗證,已關閉"', allow_blank=True)
ws.add_data_validation(dv3)
ws.add_data_validation(dv4)
dv3.add("D2:D50")
dv4.add("J2:J50")
ws.conditional_formatting.add("A2:L50", FormulaRule(formula=['AND($D2="高",$J2<>"已關閉")'], fill=RED_FILL))
ws.cell(row=52, column=1, value="嚴重度：高 = 影響金額或機率，不可上線；中 = 功能異常但不影響金額；低 = 文字、排版等。").font = F_NOTE
ws.freeze_panes = "B2"

# ---------- 回歸測試 ----------
ws = wb.create_sheet("回歸測試")
title(ws, "BUG-01 回歸測試（v1.3.1）", "修正免費遊戲倍數後，重測原本失敗的案例，以及所有可能受影響的免費遊戲相關案例")
header(ws, 4, ["案例", "測試項目", "結果（引用測試案例）", "重測原因"], [10, 30, 20, 46])
reg = [("TC-003", "分散賠付與免費遊戲設定一致", "直接對應錯誤的設定值"), ("TC-017", "免費場次", "同屬免費遊戲設定"),
       ("TC-018", "免費遊戲贏分倍數", "原本失敗的案例"), ("TC-019", "免費遊戲中再觸發", "免費遊戲流程可能受影響"),
       ("TC-020", "免費遊戲中分散賠付也乘倍數", "倍數套用範圍"), ("TC-030", "RTP 落在理論值 ±3 標準誤", "倍數直接影響 RTP"),
       ("TC-031", "免費遊戲觸發率與平均場次", "確認修正沒有影響觸發邏輯")]
for i, (tc, name, why) in enumerate(reg, 5):
    put(ws, i, 1, tc)
    put(ws, i, 2, name)
    put(ws, i, 3, f'=INDEX(測試案例!$H$2:$H$200,MATCH(A{i},測試案例!$A$2:$A$200,0))', F_LINK)
    put(ws, i, 4, why)
RL = 4 + len(reg)
ws.conditional_formatting.add(f"C5:C{RL}", FormulaRule(formula=['C5="通過"'], fill=GREEN_FILL))
ws.conditional_formatting.add(f"C5:C{RL}", FormulaRule(formula=['C5<>"通過"'], fill=RED_FILL))
put(ws, RL + 2, 1, "回歸結果", F_BOLD, fill=FILL_SUB)
put(ws, RL + 2, 2, f'=IF(COUNTIF(C5:C{RL},"通過")={len(reg)},"全部通過","有未通過項目")', F_BOLD, fill=FILL_SUB)

for s in wb.worksheets:
    s.sheet_view.showGridLines = s.title not in ("摘要", "回歸測試", "數學驗收")
wb.save(os.path.join(HERE, "金鯉躍門_v1.3.1_驗收.xlsx"))
print("ok", {k: (round(v["win"] / v["n"] / 20, 4), v["trig"]) for k, v in sims.items()})
